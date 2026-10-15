"""
Independent validation of every value drawn in the eCRC figures and quoted in the deck.

Reads the source workbooks with openpyxl (no shared code with ecrc_maps.py), recomputes the
expected values, and compares them to what ecrc_maps.py logged at draw time (ECRC_LOG json):
values, color class read back from the drawn polygons, hatching / no-data status, labels,
tick labels, bar heights, line y-values. Also checks numbers quoted in slide text and notes.

Usage: python validate_figures.py <summary.xlsx> <by-month.xlsx> <fig_log.json> <deck.pptx | -> <out.csv>
"""
import sys, json, re, csv, math, zipfile
import openpyxl

SUMMARY, MONTHLY, LOGF, DECK, OUT = sys.argv[1:6]
L = json.load(open(LOGF))
rows = []
def check(fig, element, key, expected, drawn, ok, note=""):
    rows.append(dict(figure=fig, element=element, key=key, expected=expected, drawn=drawn,
                     status="PASS" if ok else "FAIL", note=note))
def warn(fig, element, key, expected, drawn, note):
    rows.append(dict(figure=fig, element=element, key=key, expected=expected, drawn=drawn, status="WARN", note=note))
close = lambda a, b, t=1e-9: (a is None and b is None) or (a is not None and b is not None and abs(a - b) <= t)
AREA = {"5": "5", "6": "6", "7": "7", "8_1": "81", "8_2": "82", "9": "9", "10": "10", "11": "11", "12": "12", "13": "13"}
LAB = {"81": "8-1", "82": "8-2"}
lab = lambda c: LAB.get(c, c)
pct0 = lambda v: f"{round(v * 100 + 1e-9):.0f}%"

# ---------------- source: summary ----------------
ws = openpyxl.load_workbook(SUMMARY, data_only=True).active
grid = [[c.value for c in r] for r in ws.iter_rows()]
hr = next(i for i, r in enumerate(grid) if r[1] == "Catch area")
hdr = [h for h in grid[hr][1:] if h is not None]
S, TOTAL = {}, None
for r in grid[hr + 1:]:
    if r[1] is None: continue
    rec = dict(zip(hdr, r[1:1 + len(hdr)]))
    if str(r[1]) == "Total": TOTAL = rec; break
    S[AREA[str(r[1])]] = rec
OVERALL = {r[1]: r[3] for r in grid[:hr] if r[1] and r[3] is not None}
AREAS = list(S)

# ---------------- source: monthly ----------------
wm = openpyxl.load_workbook(MONTHLY, data_only=True).active
mg = [[c.value for c in r] for r in wm.iter_rows()]
def block(title):
    i = next(k for k, r in enumerate(mg) if isinstance(r[0], str) and r[0].startswith(title))
    h = [x for x in mg[i + 1] if x is not None]
    out, allrow = {}, None
    for r in mg[i + 2:]:
        if r[0] is None: break
        rec = dict(zip(h, r[:len(h)]))
        if str(r[0]) == "All areas": allrow = rec; break
        out[AREA[str(r[0])]] = rec
    return h, out, allrow
num = lambda x: None if x in (None, "–", "-") else float(x)
mh, NG, NALL = block("eCRC interviews by area and month")
_, FG, FALL = block("% of fish via eCRC")
_, RG, RALL = block("% anglers refused scan")
MONTHS = [m for m in mh if str(m).startswith("2026-")]
ai = next(k for k, r in enumerate(mg) if r[0] == "Month")
AM = []
for r in mg[ai + 1:]:
    if r[0] in (None, "Total"): break
    AM.append(dict(zip([x for x in mg[ai] if x], r)))
AMTOT = dict(zip([x for x in mg[ai] if x], mg[ai + 1 + len(AM)]))

# ---------------- source consistency (data, not figures) ----------------
F = "source data"
for col in ["Total interviews", "eCRC interviews", "Barcode scans", "Interviews w/ a refusal", "eCRC anglers",
            "Refusing anglers", "Total fish", "eCRC fish"]:
    sm = sum(float(S[c][col]) for c in AREAS)
    check(F, "summary Total row = sum of areas", col, round(sm, 3), round(float(TOTAL[col]), 3), close(sm, float(TOTAL[col]), 0.01))
for c in AREAS:
    check(F, "% fish = eCRC fish / Total fish", lab(c), round(S[c]["eCRC fish"] / S[c]["Total fish"], 6),
          round(S[c]["% fish via eCRC"], 6), close(S[c]["eCRC fish"] / S[c]["Total fish"], S[c]["% fish via eCRC"], 1e-6))
    check(F, "% refused = refusing / eCRC anglers", lab(c), round(S[c]["Refusing anglers"] / S[c]["eCRC anglers"], 6),
          round(S[c]["% anglers refused"], 6), close(S[c]["Refusing anglers"] / S[c]["eCRC anglers"], S[c]["% anglers refused"], 1e-6))
    check(F, "summary eCRC interviews = monthly grid season", lab(c), S[c]["eCRC interviews"], NG[c]["Season"],
          S[c]["eCRC interviews"] == NG[c]["Season"])
    check(F, "summary % refused = monthly grid season", lab(c), round(S[c]["% anglers refused"], 6), round(num(RG[c]["Season"]), 6),
          close(S[c]["% anglers refused"], num(RG[c]["Season"]), 1e-6))
    d = abs(S[c]["% fish via eCRC"] - num(FG[c]["Season"]))
    (check if d < 0.005 else warn)(F, "summary % fish vs monthly grid season", lab(c), round(S[c]["% fish via eCRC"], 4),
          round(num(FG[c]["Season"]), 4), *(([True, f"diff {d:.4f}; summary pools licenses over season"]) if d < 0.005 else [f"diff {d:.4f}"]))
for k_over, k_tab in [("Fish reported via eCRC", "eCRC fish"), ("% fish via eCRC", "% fish via eCRC"),
                      ("Interviews w/ eCRC harvest", "eCRC interviews"), ("Total fish", "Total fish"),
                      ("Interviews w/ harvest", "Total interviews")]:
    a, b = float(OVERALL[k_over]), float(TOTAL[k_tab])
    if close(a, b, 0.01): check(F, "Overall block = table Total row", k_over, a, b, True)
    else: warn(F, "Overall block = table Total row", k_over, round(b, 4), round(a, 4), "Overall block not updated to the license-proportion estimate; figures use the table")
for m in MONTHS:
    sm = sum(NG[c][m] for c in AREAS)
    check(F, "monthly n grid: sum of areas = All areas", m, sm, NALL[m], sm == NALL[m])
    am = next(r for r in AM if r["Month"] == m)
    check(F, "monthly n grid All areas = all-areas table", m, am["eCRC interviews"], NALL[m], am["eCRC interviews"] == NALL[m])
    check(F, "all-areas % refused = refusing / eCRC anglers", m, round(am["Refusing anglers"] / am["eCRC anglers"], 6),
          round(am["% anglers refused"], 6), close(am["Refusing anglers"] / am["eCRC anglers"], am["% anglers refused"], 1e-6))
    check(F, "all-areas % fish = est. eCRC fish / total fish", m, round(am["eCRC fish (est.)"] / am["Total fish"], 6),
          round(am["% fish via eCRC (est.)"], 6), close(am["eCRC fish (est.)"] / am["Total fish"], am["% fish via eCRC (est.)"], 1e-6))

# ---------------- maps 01, 02 ----------------
def cls(v, br): return sum(v >= b for b in br)
for fig, col, br in [("01_season_pct_fish_ecrc.png", "% fish via eCRC", [.15, .20, .25, .30]),
                     ("02_season_refusal_rate.png", "% anglers refused", [.05, .10, .20, .30])]:
    D, LB = L[fig], L[fig + ":labels"]
    for c in AREAS:
        v = float(S[c][col]); d = D[c]
        check(fig, "value", lab(c), round(v, 6), round(d["v"], 6), close(v, d["v"]))
        check(fig, "color class (read from drawn polygon)", lab(c), cls(v, br), d["class"], cls(v, br) == d["class"],
              f"breaks {br}")
        check(fig, "hatched (n < 10)", lab(c), S[c]["eCRC interviews"] < 10, d["hatched"], (S[c]["eCRC interviews"] < 10) == d["hatched"])
        exp = f"{lab(c)}\n{pct0(v)}"
        check(fig, "label", lab(c), exp.replace("\n", " "), LB[c].replace("\n", " "), exp == LB[c])

# ---------------- map 03 ----------------
fig = "03_season_ecrc_sample_size.png"; C3 = L["03:circles"]
k = [C3[c]["outer_size"] / S[c]["eCRC interviews"] for c in AREAS]
for c in AREAS:
    n, sc = int(S[c]["eCRC interviews"]), int(S[c]["Barcode scans"])
    check(fig, "label: n", lab(c), f"($\\it{{n}}$ = {n:,})", C3[c]["labels"][1], C3[c]["labels"][1] == f"($\\it{{n}}$ = {n:,})")
    check(fig, "label: scans", lab(c), f"{sc:,} scans", C3[c]["labels"][2], C3[c]["labels"][2] == f"{sc:,} scans")
    check(fig, "circle area proportional to n (same scale all areas)", lab(c), round(k[0], 6),
          round(C3[c]["outer_size"] / n, 6), close(k[0], C3[c]["outer_size"] / n, 1e-6))
    check(fig, "inner/outer area = scans/n", lab(c), round(sc / n, 6), round(C3[c]["inner_size"] / C3[c]["outer_size"], 6),
          close(sc / n, C3[c]["inner_size"] / C3[c]["outer_size"], 1e-6))

# ---------------- maps 04, 05 (monthly) ----------------
for fig, G, br in [("04_monthly_map_pct_fish_ecrc.png", FG, [.15, .20, .25, .30]),
                   ("05_monthly_map_refusal.png", RG, [.05, .10, .20, .30])]:
    for m in MONTHS:
        D = L[f"{fig}:{m}"]
        for c in AREAS:
            n = NG[c][m]; v = num(G[c][m]); d = D[c]
            exp_v = None if (n == 0 or v is None) else v
            key = f"{m} {lab(c)}"
            check(fig, "value", key, exp_v, d["v"], close(exp_v, d["v"]))
            check(fig, "n", key, n, d["n"], n == d["n"])
            check(fig, "no-data (n = 0 or '–')", key, exp_v is None, d["nodata"], (exp_v is None) == d["nodata"])
            if exp_v is not None:
                check(fig, "color class (read from drawn polygon)", key, cls(exp_v, br), d["class"], cls(exp_v, br) == d["class"])
                check(fig, "hatched (n < 10)", key, n < 10, d["hatched"], (n < 10) == d["hatched"])
            elif v is not None and n == 0:
                warn(fig, "value present with n = 0", key, "no data", v, "source has a value where n = 0; drawn as no data")
        t = L[f"{fig}:titles"][m]; tn = int(re.search(r"= ([\d,]+)\)", t).group(1).replace(",", ""))
        check(fig, "panel title n = All areas", m, NALL[m], tn, NALL[m] == tn)

# ---------------- plot 06 ----------------
fig = "06_monthly_lines_all_areas.png"
for key, col in [("pct_fish", "% fish via eCRC (est.)"), ("pct_refused", "% anglers refused")]:
    for i, r in enumerate(AM):
        check(fig, f"line y: {key}", r["Month"], round(r[col], 6), round(L["06:lines"][key][i], 6), close(r[col], L["06:lines"][key][i]))
for i, r in enumerate(AM):
    tn = int(re.search(r"= ([\d,]+)\)", L["06:xticks"][i]).group(1).replace(",", ""))
    check(fig, "x label n", r["Month"], r["eCRC interviews"], tn, r["eCRC interviews"] == tn)

# ---------------- plot 07 ----------------
fig = "07_monthly_lines_by_area.png"; P = L["07:panels"]
for c in AREAS:
    sn = int(re.search(r"= ([\d,]+)\)", P[c]["strip"]).group(1).replace(",", ""))
    check(fig, "strip n = season eCRC interviews", lab(c), int(S[c]["eCRC interviews"]), sn, sn == S[c]["eCRC interviews"])
    for key, G in [("pct_fish", FG), ("pct_refused", RG)]:
        for i, m in enumerate(MONTHS):
            n, v = NG[c][m], num(G[c][m])
            es = v if (v is not None and n >= 10) else None
            eh = v if (v is not None and n < 10) else None
            k2 = f"{m} {lab(c)} {key}"
            check(fig, "solid point (n >= 10)", k2, es, P[c]["solid"][key][i], close(es, P[c]["solid"][key][i]))
            check(fig, "hollow point (n < 10)", k2, eh, P[c]["hollow"][key][i], close(eh, P[c]["hollow"][key][i]))

# ---------------- plot 08 ----------------
fig = "08_scans_refusals_by_area.png"; B = L["08:bars"]; labs = L["08:labels"]
for i, c in enumerate(AREAS):
    n = S[c]["eCRC interviews"]
    check(fig, "bar: scans per eCRC interview", lab(c), round(S[c]["Barcode scans"] / n, 6), round(B[c]["scan_h"], 6),
          close(S[c]["Barcode scans"] / n, B[c]["scan_h"]))
    check(fig, "bar: refusal interviews per eCRC interview", lab(c), round(S[c]["Interviews w/ a refusal"] / n, 6),
          round(B[c]["ref_h"], 6), close(S[c]["Interviews w/ a refusal"] / n, B[c]["ref_h"]))
    check(fig, "count label: scans", lab(c), f"{int(S[c]['Barcode scans']):,}", labs[i], labs[i] == f"{int(S[c]['Barcode scans']):,}")
    check(fig, "count label: refusals", lab(c), f"{int(S[c]['Interviews w/ a refusal']):,}", labs[len(AREAS) + i],
          labs[len(AREAS) + i] == f"{int(S[c]['Interviews w/ a refusal']):,}")
    tn = int(re.search(r"= ([\d,]+)\)", L["08:xticks"][i]).group(1).replace(",", ""))
    check(fig, "x label n", lab(c), int(n), tn, tn == n)
note = L["captions"][fig]["note"]
ts, tn = int(TOTAL["Barcode scans"]), int(TOTAL["eCRC interviews"])
exp = f"{ts:,} scans in {tn:,} eCRC interviews ({pct0(ts / tn)})"
check(fig, "caption overall", "all areas", exp, "present" if exp in note else "missing", exp in note)

if DECK not in ("-", "", "none"):   # deck checks are optional; pass "-" to skip
    # ---------------- deck text (slides and notes) ----------------
    z = zipfile.ZipFile(DECK)
    order = re.findall(r'r:id="(rId\d+)"', re.search(r"<p:sldIdLst>.*?</p:sldIdLst>", z.read("ppt/presentation.xml").decode()).group(0))
    rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="(slides/slide\d+\.xml)"', z.read("ppt/_rels/presentation.xml.rels").decode()))
    def text_of(path):
        x = z.read(path).decode()
        return " ".join("".join(re.findall(r"<a:t[^>]*>([^<]*)</a:t>", p)) for p in re.findall(r"<a:p>.*?</a:p>", x, flags=re.S)).replace("&amp;", "&")
    SL = {}
    for i, rid in enumerate(order, 1):
        sp = "ppt/" + rels[rid]; t = text_of(sp)
        nrel = "ppt/slides/_rels/" + sp.split("/")[-1] + ".rels"
        nm = re.search(r'Target="\.\./notesSlides/(notesSlide\d+\.xml)"', z.read(nrel).decode())
        SL[i] = t + " || NOTES: " + (text_of("ppt/notesSlides/" + nm.group(1)) if nm else "")

    def find_slide(title):
        return next(i for i, t in SL.items() if t.startswith(title))
    tot = lambda col: float(TOTAL[col])
    claims = []
    s6 = find_slide("Early results")
    claims += [(s6, f"{int(tot('Total interviews')):,} harvest interviews"),
               (s6, f"{int(OVERALL['Total anglers']):,} anglers and {int(tot('Total fish')):,} fish"),
               (s6, f"{int(tot('eCRC interviews')):,} interviews ({pct0(tot('eCRC interviews') / tot('Total interviews'))})"),
               (s6, f"About {pct0(tot('eCRC fish') / tot('Total fish'))} of fish"),
               (s6, f"{pct0(tot('Refusing anglers') / tot('eCRC anglers'))} of eCRC anglers"),
               (s6, f"({int(tot('Refusing anglers')):,} of {int(tot('eCRC anglers')):,})"),
               (s6, f"{int(tot('Barcode scans')):,} barcode scans – {round(100 * tot('Barcode scans') / tot('eCRC interviews'))} per 100"),
               (s6, f"{pct0(sum(S[c]['eCRC interviews'] for c in ['82', '9', '10']) / tot('eCRC interviews'))} of eCRC interviews – from Areas 8-2, 9 and 10")]
    s7 = find_slide("Where the eCRC sample")
    top3 = sorted(AREAS, key=lambda c: -S[c]["Barcode scans"])[:3]
    pre_july = [c for c in AREAS if any(NG[c][m] > 0 for m in MONTHS[:3])]
    claims += [(s7, f"{int(S['10']['eCRC interviews']):,}, {int(S['82']['eCRC interviews']):,} and {int(S['9']['eCRC interviews']):,} eCRC interviews"),
               (s7, f"most in Areas {', '.join(lab(c) for c in top3[:2])} and {lab(top3[2])} ({', '.join(str(int(S[c]['Barcode scans'])) for c in top3)})"),
               (s7, f"only Areas {', '.join(lab(c) for c in pre_july[:-1])} and {lab(pre_july[-1])} had eCRC interviews")]
    lo = [S[c]["eCRC interviews"] for c in ["81", "13"]]
    check("deck", "slide text", f"slide {s7}: Areas 8-1 and 13 fewer than 40", "max < 40", max(lo), max(lo) < 40)
    r567 = [int(S[c]["eCRC interviews"]) for c in ["5", "6", "7"]]
    claims += [(s7, f"Areas 5, 6 and 7 – {min(r567)} to {max(r567)} each")]
    s8 = find_slide("Estimated share")
    claims += [(s8, f"{pct0(tot('eCRC fish') / tot('Total fish'))} overall"),
               (s8, f"of {int(tot('Total fish')):,} fish"),
               (s8, f"Area 10 ({pct0(S['10']['% fish via eCRC'])}), 9 ({pct0(S['9']['% fish via eCRC'])})"),
               (s8, f"8-2 and 7 ({pct0(S['82']['% fish via eCRC'])})"),
               (s8, f"Strait {pct0(min(S[c]['% fish via eCRC'] for c in ['5','6']))[:-1]}–{pct0(max(S[c]['% fish via eCRC'] for c in ['5','6']))}"),
               (s8, f"South Sound and Hood Canal {pct0(min(S[c]['% fish via eCRC'] for c in ['11','12','13']))[:-1]}–{pct0(max(S[c]['% fish via eCRC'] for c in ['11','12','13']))}"),
               (s8, f"8-1 {pct0(S['81']['% fish via eCRC'])}")]
    if pct0(S["82"]["% fish via eCRC"]) != pct0(S["7"]["% fish via eCRC"]):
        check("deck", "slide text", f"slide {s8}: 8-2 and 7 share one value", pct0(S["82"]["% fish via eCRC"]), pct0(S["7"]["% fish via eCRC"]), False)
    a10 = [num(FG["10"][m]) for m in MONTHS[2:]]
    claims += [(s8, f"Area 10 – {pct0(min(a10))[:-1]}–{pct0(max(a10))} every month from June through October")]
    s9 = find_slide("Barcode scan refusals")
    a6 = [num(RG["6"][m]) for m in MONTHS if num(RG["6"][m]) is not None]
    others = [S[c]["% anglers refused"] for c in ["5", "7", "81", "82", "9", "10", "11"]]
    claims += [(s9, f"{pct0(tot('Refusing anglers') / tot('eCRC anglers'))} overall – {int(tot('Refusing anglers')):,} of {int(tot('eCRC anglers')):,}"),
               (s9, f"Area 6 – {pct0(S['6']['% anglers refused'])} season to date, between {pct0(min(a6))} and {pct0(max(a6))} every month"),
               (s9, f"Area 13 – {pct0(S['13']['% anglers refused'])}, but from only {int(S['13']['eCRC anglers'])} eCRC anglers"),
               (s9, f"Area 12 – lowest at {pct0(S['12']['% anglers refused'])} ({int(S['12']['Refusing anglers'])} of {int(S['12']['eCRC anglers'])})"),
               (s9, f"Most other areas – {pct0(min(others))[:-1]} to {pct0(max(others))}")]
    s10 = find_slide("Barcode scans by catch area")
    sr = {c: S[c]["Barcode scans"] / S[c]["eCRC interviews"] for c in AREAS}
    hi = sorted(AREAS, key=lambda c: -sr[c])[:2]; low2 = sorted(AREAS, key=lambda c: sr[c])[:2]
    claims += [(s10, f"Overall {int(tot('Barcode scans')):,} scans in {int(tot('eCRC interviews')):,} eCRC interviews, about {round(100 * tot('Barcode scans') / tot('eCRC interviews'))} per 100"),
               (s10, f"Highest in Areas {lab(hi[0])} ({round(100 * sr[hi[0]])}) and {lab(hi[1])} ({round(100 * sr[hi[1]])})"),
               (s10, f"lowest in Area {lab(low2[0])} ({round(100 * sr[low2[0]])}) and {lab(low2[1])} ({round(100 * sr[low2[1]])})"),
               (s10, f"Area 6 has more interviews with a refusal ({int(S['6']['Interviews w/ a refusal'])}) than scans ({int(S['6']['Barcode scans'])})")]
    s11 = find_slide("Season trend")
    mrec = {r["Month"]: r for r in AM}
    fj = [mrec[m]["% fish via eCRC (est.)"] for m in ["2026-07", "2026-08", "2026-09"]]
    claims += [(s11, f"Refusals: {pct0(mrec['2026-04']['% anglers refused'])} in April, {pct0(mrec['2026-09']['% anglers refused'])} in September, back up to {pct0(mrec['2026-10']['% anglers refused'])} in early October (partial month, {int(mrec['2026-10']['eCRC anglers'])} eCRC anglers)"),
               (s11, f"eCRC fish share: {pct0(min(fj))[:-1]}-{pct0(max(fj))} July through September"),
               (s11, f"May has {int(mrec['2026-05']['eCRC interviews'])} eCRC interviews")]
    for s, txt in claims:
        check("deck", "slide text / notes", f"slide {s}", txt, "present" if txt in SL[s] else "missing", txt in SL[s])

with open(OUT, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
from collections import Counter
byfig = {}
for r in rows: byfig.setdefault(r["figure"], Counter())[r["status"]] += 1
for k, v in byfig.items(): print(f"{k:38s} " + "  ".join(f"{s}={v[s]}" for s in ["PASS", "FAIL", "WARN"]))
print("TOTAL", Counter(r["status"] for r in rows))
for r in rows:
    if r["status"] != "PASS": print(r["status"], "|", r["figure"], "|", r["element"], "|", r["key"], "| expected:", r["expected"], "| drawn:", r["drawn"], "|", r["note"])
